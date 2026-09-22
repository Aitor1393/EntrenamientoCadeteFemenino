from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
import datetime as dt

# ---------- estilos ----------
F = "Arial"
titulo = Font(name=F, size=14, bold=True, color="1F3864")
subtitulo = Font(name=F, size=11, bold=True, color="1F3864")
cab = Font(name=F, size=9, bold=True, color="FFFFFF")
normal = Font(name=F, size=10)
entrada = Font(name=F, size=10, color="0000FF")
calc = Font(name=F, size=10, color="000000")
bold = Font(name=F, size=10, bold=True)
peque = Font(name=F, size=9, italic=True, color="595959")

fill_cab = PatternFill("solid", fgColor="1F3864")
fill_in = PatternFill("solid", fgColor="FFF2CC")
fill_calc = PatternFill("solid", fgColor="F2F2F2")
fill_sec = PatternFill("solid", fgColor="D9E2F3")

verde = PatternFill("solid", fgColor="C6EFCE")
rojo = PatternFill("solid", fgColor="FFC7CE")

thin = Side(style="thin", color="BFBFBF")
borde = Border(left=thin, right=thin, top=thin, bottom=thin)

ctr = Alignment(horizontal="center", vertical="center")
ctr_wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()

# =========================================================
# HOJA 1 — INSTRUCCIONES
# =========================================================
ws = wb.active
ws.title = "Instrucciones"
ws.sheet_view.showGridLines = False

filas = [
    ("SEGUIMIENTO DE TEMPORADA 2026/2027", titulo, None),
    ("Cadete Femenino Preferente · Serie B · martes 75 min y viernes 105 min", subtitulo, None),
    ("", None, None),
    ("CÓMO SE USA", subtitulo, None),
    ("Solo se escribe en las celdas amarillas con texto azul. El resto son fórmulas y se calculan solas.", normal, None),
    ("", None, None),
    ("Hoja «Partidos»", bold, None),
    ("Cinco datos son el mínimo imprescindible y se pueden contar desde el banquillo:", normal, None),
    ("   1.  Rebotes defensivos capturados por nosotras", normal, None),
    ("   2.  Rebotes ofensivos que captura el rival  (los dos juntos dan el % de rebote defensivo)", normal, None),
    ("   3.  Rebotes ofensivos capturados por nosotras", normal, None),
    ("   4.  Puntos anotados en los primeros 8 segundos de posesión", normal, None),
    ("   5.  Pérdidas", normal, None),
    ("Los puntos a favor y en contra salen del acta. Los rebotes defensivos del rival son opcionales:", normal, None),
    ("sin ellos no se calcula el % de rebote ofensivo, pero el resto de la hoja funciona igual.", normal, None),
    ("", None, None),
    ("Hoja «Minutos»", bold, None),
    ("Minutos jugados por cada jugadora. La columna «Partidos < 10 min» vigila el criterio de reparto:", normal, None),
    ("durante M1 y M2 nadie debería bajar de 10 minutos. En M3 (J19 a J22) ese criterio se relaja.", normal, None),
    ("La fila «Suma» debe rondar los 200 minutos por partido: sirve para detectar errores de anotación.", normal, None),
    ("", None, None),
    ("Hojas «Test físico» y «Test técnico»", bold, None),
    ("Cuatro mediciones: septiembre, enero, marzo y mayo. La columna «Mejora» se calcula sola y ya", normal, None),
    ("tiene en cuenta si la prueba mejora subiendo (salto, tiros) o bajando (sprint, agilidad, manejo).", normal, None),
    ("", None, None),
    ("Hoja «Panel»", bold, None),
    ("Resumen automático. No se escribe nada. Se revisa al cerrar cada bloque: diciembre, marzo y abril.", normal, None),
    ("", None, None),
    ("OBJETIVOS DE RENDIMIENTO", subtitulo, None),
]
r = 1
for txt, fnt, _ in filas:
    c = ws.cell(row=r, column=1, value=txt)
    if fnt:
        c.font = fnt
    r += 1

obj = [
    ("Indicador", "Objetivo", "Por qué"),
    ("% rebote defensivo", "> 70 %", "Es el arranque del modelo de juego"),
    ("Puntos en transición", "> 20 por partido", "Mide si el contraataque existe de verdad"),
    ("% rebote ofensivo", "> 30 %", "Segundas opciones"),
    ("Pérdidas", "< 18 por partido", "Correr sin control se convierte en pérdidas"),
]
r += 1
inicio_obj = r
for i, fila in enumerate(obj):
    for j, v in enumerate(fila):
        c = ws.cell(row=r, column=1 + j, value=v)
        c.border = borde
        if i == 0:
            c.font = cab
            c.fill = fill_cab
            c.alignment = ctr
        else:
            c.font = normal
            if j == 1:
                c.alignment = ctr
    r += 1

r += 1
ws.cell(row=r, column=1, value="LEYENDA DE COLORES").font = subtitulo
r += 1
leg = [("Celda a rellenar", fill_in, entrada), ("Celda calculada", fill_calc, calc)]
for txt, fl, fn in leg:
    c = ws.cell(row=r, column=1, value=txt)
    c.fill = fl
    c.font = fn
    c.border = borde
    r += 1

r += 2
ws.cell(row=r, column=1, value="Las fechas asumen que se juega en sábado. Confírmalas con el calendario oficial de la FBM tras el sorteo.").font = peque

ws.column_dimensions["A"].width = 62
ws.column_dimensions["B"].width = 20
ws.column_dimensions["C"].width = 48

# =========================================================
# HOJA 2 — PARTIDOS
# =========================================================
ws = wb.create_sheet("Partidos")
ws.sheet_view.showGridLines = False

ws["A1"] = "REGISTRO DE PARTIDOS"
ws["A1"].font = titulo
ws["A2"] = "Rellena solo las celdas amarillas. Fila 5 = ejemplo, bórrala cuando empieces."
ws["A2"].font = peque

cabeceras = [
    ("Jornada", 10), ("Fecha", 11), ("Rival", 22), ("C/F", 6),
    ("PF", 6), ("PC", 6), ("Dif.", 7), ("V/D", 6),
    ("RD\nnuestros", 10), ("RO\nrival", 9), ("% Reb\nDef", 9),
    ("RO\nnuestros", 10), ("RD\nrival", 9), ("% Reb\nOf", 9),
    ("Pts\ntransición", 11), ("Pérdidas", 10), ("Notas", 30),
]
HDR = 4
for j, (txt, w) in enumerate(cabeceras, start=1):
    c = ws.cell(row=HDR, column=j, value=txt)
    c.font = cab
    c.fill = fill_cab
    c.alignment = ctr_wrap
    c.border = borde
    ws.column_dimensions[get_column_letter(j)].width = w
ws.row_dimensions[HDR].height = 30

jornadas = [
    ("J1", dt.date(2026, 10, 3)), ("J2", dt.date(2026, 10, 17)), ("J3", dt.date(2026, 10, 24)),
    ("J4", dt.date(2026, 11, 7)), ("J5", dt.date(2026, 11, 14)), ("J6", dt.date(2026, 11, 21)),
    ("J7", dt.date(2026, 11, 28)), ("J8", dt.date(2026, 12, 12)), ("J9", dt.date(2026, 12, 19)),
    ("J10", dt.date(2027, 1, 16)), ("J11", dt.date(2027, 1, 23)), ("J12", dt.date(2027, 1, 30)),
    ("J13", dt.date(2027, 2, 6)), ("J14", dt.date(2027, 2, 13)), ("J15", dt.date(2027, 2, 20)),
    ("J16", dt.date(2027, 2, 27)), ("J17", dt.date(2027, 3, 6)), ("J18", dt.date(2027, 3, 13)),
    ("J19", dt.date(2027, 4, 3)), ("J20", dt.date(2027, 4, 10)), ("J21", dt.date(2027, 4, 17)),
    ("J22", dt.date(2027, 4, 24)),
    ("2ªF-1", dt.date(2027, 5, 8)), ("2ªF-2", dt.date(2027, 5, 15)),
    ("2ªF-3", dt.date(2027, 5, 22)), ("2ªF-4", dt.date(2027, 5, 29)),
]

R0 = HDR + 1                      # 5
R1 = R0 + len(jornadas) - 1       # 30
cols_in = [3, 4, 5, 6, 9, 10, 12, 13, 15, 16, 17]

for i, (jn, fecha) in enumerate(jornadas):
    r = R0 + i
    ws.cell(row=r, column=1, value=jn).font = bold
    ws.cell(row=r, column=1).alignment = ctr
    fc = ws.cell(row=r, column=2, value=fecha)
    fc.number_format = "DD/MM/YYYY"
    fc.font = normal
    fc.alignment = ctr
    for j in cols_in:
        c = ws.cell(row=r, column=j)
        c.font = entrada
        c.fill = fill_in
        if j not in (3, 17):
            c.alignment = ctr
    ws.cell(row=r, column=7, value=f"=IF(E{r}=\"\",\"\",E{r}-F{r})")
    ws.cell(row=r, column=8, value=f"=IF(E{r}=\"\",\"\",IF(E{r}>F{r},\"V\",\"D\"))")
    ws.cell(row=r, column=11, value=f"=IFERROR(I{r}/(I{r}+J{r}),\"\")")
    ws.cell(row=r, column=14, value=f"=IFERROR(L{r}/(L{r}+M{r}),\"\")")
    for j in (7, 8, 11, 14):
        c = ws.cell(row=r, column=j)
        c.font = calc
        c.fill = fill_calc
        c.alignment = ctr
    ws.cell(row=r, column=11).number_format = "0.0%"
    ws.cell(row=r, column=14).number_format = "0.0%"
    for j in range(1, 18):
        ws.cell(row=r, column=j).border = borde
    if jn.startswith("2ªF"):
        ws.cell(row=r, column=1).fill = fill_sec

# fila de ejemplo
ejemplo = {3: "Ejemplo CB", 4: "C", 5: 54, 6: 48, 9: 24, 10: 9, 12: 11, 13: 22, 15: 22, 16: 16,
           17: "Buen rebote, muchas pérdidas en el 3er cuarto"}
for j, v in ejemplo.items():
    ws.cell(row=R0, column=j, value=v)

# medias
def bloque_media(r, etiqueta, ini, fin):
    ws.cell(row=r, column=1, value=etiqueta).font = bold
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    ws.cell(row=r, column=1).alignment = Alignment(horizontal="right", vertical="center")
    for j in (5, 6, 7, 9, 10, 11, 12, 13, 14, 15, 16):
        col = get_column_letter(j)
        ws.cell(row=r, column=j, value=f"=IFERROR(AVERAGE({col}{ini}:{col}{fin}),\"\")")
        c = ws.cell(row=r, column=j)
        c.font = bold
        c.fill = fill_sec
        c.alignment = ctr
        c.border = borde
        c.number_format = "0.0%" if j in (11, 14) else "0.0"
    ws.cell(row=r, column=8, value=f"=IFERROR(COUNTIF(H{ini}:H{fin},\"V\")&\"-\"&COUNTIF(H{ini}:H{fin},\"D\"),\"\")")
    ws.cell(row=r, column=8).font = bold
    ws.cell(row=r, column=8).fill = fill_sec
    ws.cell(row=r, column=8).alignment = ctr
    ws.cell(row=r, column=8).border = borde
    for j in range(1, 5):
        ws.cell(row=r, column=j).fill = fill_sec
        ws.cell(row=r, column=j).border = borde

R_M1 = R1 + 2
R_M2 = R1 + 3
bloque_media(R_M1, "MEDIA 1ª FASE  (J1-J22)   ", R0, R0 + 21)
bloque_media(R_M2, "MEDIA TEMPORADA   ", R0, R1)

# formato condicional
rng_rd = f"K{R0}:K{R1}"
ws.conditional_formatting.add(rng_rd, FormulaRule(formula=[f'AND(K{R0}<>"",K{R0}>=0.7)'], fill=verde))
ws.conditional_formatting.add(rng_rd, FormulaRule(formula=[f'AND(K{R0}<>"",K{R0}<0.7)'], fill=rojo))
rng_ro = f"N{R0}:N{R1}"
ws.conditional_formatting.add(rng_ro, FormulaRule(formula=[f'AND(N{R0}<>"",N{R0}>=0.3)'], fill=verde))
ws.conditional_formatting.add(rng_ro, FormulaRule(formula=[f'AND(N{R0}<>"",N{R0}<0.3)'], fill=rojo))
rng_tr = f"O{R0}:O{R1}"
ws.conditional_formatting.add(rng_tr, FormulaRule(formula=[f'AND(O{R0}<>"",O{R0}>=20)'], fill=verde))
ws.conditional_formatting.add(rng_tr, FormulaRule(formula=[f'AND(O{R0}<>"",O{R0}<20)'], fill=rojo))
rng_pe = f"P{R0}:P{R1}"
ws.conditional_formatting.add(rng_pe, FormulaRule(formula=[f'AND(P{R0}<>"",P{R0}<=18)'], fill=verde))
ws.conditional_formatting.add(rng_pe, FormulaRule(formula=[f'AND(P{R0}<>"",P{R0}>18)'], fill=rojo))

ws.freeze_panes = "C5"

nota = ws.cell(row=R_M2 + 2, column=1,
               value="RD = rebote defensivo · RO = rebote ofensivo · C/F = casa o fuera. "
                     "Verde y rojo marcan si se cumple el objetivo del indicador.")
nota.font = peque

# =========================================================
# HOJA 3 — MINUTOS
# =========================================================
ws = wb.create_sheet("Minutos")
ws.sheet_view.showGridLines = False
ws["A1"] = "MINUTOS POR JUGADORA"
ws["A1"].font = titulo
ws["A2"] = "Criterio: nadie por debajo de 10 minutos durante M1 y M2 (J1 a J18). En M3 (J19-J22) la rotación se acorta."
ws["A2"].font = peque

HDR = 4
ws.cell(row=HDR, column=1, value="Nº").font = cab
ws.cell(row=HDR, column=2, value="Jugadora").font = cab
for j in (1, 2):
    ws.cell(row=HDR, column=j).fill = fill_cab
    ws.cell(row=HDR, column=j).alignment = ctr
    ws.cell(row=HDR, column=j).border = borde
ws.column_dimensions["A"].width = 5
ws.column_dimensions["B"].width = 22

etiquetas = [j for j, _ in jornadas]
for i, et in enumerate(etiquetas):
    col = 3 + i
    c = ws.cell(row=HDR, column=col, value=et)
    c.font = cab
    c.fill = fill_cab
    c.alignment = ctr
    c.border = borde
    ws.column_dimensions[get_column_letter(col)].width = 6

col_ult = 2 + len(etiquetas)          # 28
extras = ["Total", "Media", "Partidos", "< 10 min"]
for i, et in enumerate(extras):
    col = col_ult + 1 + i
    c = ws.cell(row=HDR, column=col, value=et)
    c.font = cab
    c.fill = fill_cab
    c.alignment = ctr
    c.border = borde
    ws.column_dimensions[get_column_letter(col)].width = 9
ws.row_dimensions[HDR].height = 22

P0 = HDR + 1
P1 = P0 + 11
for i in range(12):
    r = P0 + i
    ws.cell(row=r, column=1, value=i + 1).font = bold
    ws.cell(row=r, column=1).alignment = ctr
    c = ws.cell(row=r, column=2, value=f"Jugadora {i+1}")
    c.font = entrada
    c.fill = fill_in
    for j in range(3, col_ult + 1):
        cc = ws.cell(row=r, column=j)
        cc.font = entrada
        cc.fill = fill_in
        cc.alignment = ctr
    ini, fin = get_column_letter(3), get_column_letter(col_ult)
    ws.cell(row=r, column=col_ult + 1, value=f"=SUM({ini}{r}:{fin}{r})")
    ws.cell(row=r, column=col_ult + 2, value=f"=IFERROR(AVERAGE({ini}{r}:{fin}{r}),\"\")")
    ws.cell(row=r, column=col_ult + 3, value=f"=COUNT({ini}{r}:{fin}{r})")
    ws.cell(row=r, column=col_ult + 4, value=f"=COUNTIF({ini}{r}:{fin}{r},\"<10\")")
    for j in range(col_ult + 1, col_ult + 5):
        cc = ws.cell(row=r, column=j)
        cc.font = bold
        cc.fill = fill_calc
        cc.alignment = ctr
    ws.cell(row=r, column=col_ult + 2).number_format = "0.0"
    for j in range(1, col_ult + 5):
        ws.cell(row=r, column=j).border = borde

# fila de comprobación
rc = P1 + 1
ws.cell(row=rc, column=2, value="Suma (≈200)").font = bold
ws.cell(row=rc, column=2).fill = fill_sec
ws.cell(row=rc, column=1).fill = fill_sec
for j in range(3, col_ult + 1):
    col = get_column_letter(j)
    ws.cell(row=rc, column=j, value=f"=IF(SUM({col}{P0}:{col}{P1})=0,\"\",SUM({col}{P0}:{col}{P1}))")
    cc = ws.cell(row=rc, column=j)
    cc.font = bold
    cc.fill = fill_sec
    cc.alignment = ctr
for j in range(1, col_ult + 1):
    ws.cell(row=rc, column=j).border = borde

ws.conditional_formatting.add(
    f"{get_column_letter(col_ult+4)}{P0}:{get_column_letter(col_ult+4)}{P1}",
    FormulaRule(formula=[f'{get_column_letter(col_ult+4)}{P0}>0'], fill=rojo))

ws.freeze_panes = "C5"
ws.cell(row=rc + 2, column=2,
        value="Escribe los nombres reales en la columna B. La fila «Suma» debe rondar 200 minutos por partido.").font = peque

# =========================================================
# HOJAS 4 y 5 — TESTS
# =========================================================
def hoja_test(nombre, titulo_txt, pruebas, nota_txt):
    w = wb.create_sheet(nombre)
    w.sheet_view.showGridLines = False
    w["A1"] = titulo_txt
    w["A1"].font = titulo
    w["A2"] = nota_txt
    w["A2"].font = peque

    HDR1, HDR2 = 4, 5
    w.cell(row=HDR1, column=1, value="Nº")
    w.cell(row=HDR1, column=2, value="Jugadora")
    for j in (1, 2):
        w.merge_cells(start_row=HDR1, start_column=j, end_row=HDR2, end_column=j)
        c = w.cell(row=HDR1, column=j)
        c.font = cab
        c.fill = fill_cab
        c.alignment = ctr
        c.border = borde
    w.column_dimensions["A"].width = 5
    w.column_dimensions["B"].width = 22

    col = 3
    bloques = []
    for etiqueta, mayor_mejor in pruebas:
        w.merge_cells(start_row=HDR1, start_column=col, end_row=HDR1, end_column=col + 4)
        c = w.cell(row=HDR1, column=col, value=etiqueta)
        c.font = cab
        c.fill = fill_cab
        c.alignment = ctr_wrap
        for k, sub in enumerate(["Sep", "Ene", "Mar", "May", "Mejora"]):
            cc = w.cell(row=HDR2, column=col + k, value=sub)
            cc.font = cab
            cc.fill = fill_cab
            cc.alignment = ctr
            cc.border = borde
            w.column_dimensions[get_column_letter(col + k)].width = 8
        for k in range(5):
            w.cell(row=HDR1, column=col + k).border = borde
        bloques.append((col, mayor_mejor))
        col += 5
    w.row_dimensions[HDR1].height = 30

    T0 = HDR2 + 1
    for i in range(12):
        r = T0 + i
        w.cell(row=r, column=1, value=i + 1).font = bold
        w.cell(row=r, column=1).alignment = ctr
        cc = w.cell(row=r, column=2, value=f"Jugadora {i+1}")
        cc.font = entrada
        cc.fill = fill_in
        for c0, mayor in bloques:
            for k in range(4):
                x = w.cell(row=r, column=c0 + k)
                x.font = entrada
                x.fill = fill_in
                x.alignment = ctr
                x.number_format = "0.00"
            sep = get_column_letter(c0)
            may = get_column_letter(c0 + 3)
            if mayor:
                f = f"=IF(OR({sep}{r}=\"\",{may}{r}=\"\"),\"\",{may}{r}-{sep}{r})"
            else:
                f = f"=IF(OR({sep}{r}=\"\",{may}{r}=\"\"),\"\",{sep}{r}-{may}{r})"
            x = w.cell(row=r, column=c0 + 4, value=f)
            x.font = bold
            x.fill = fill_calc
            x.alignment = ctr
            x.number_format = "0.00"
        for j in range(1, col):
            w.cell(row=r, column=j).border = borde
    T1 = T0 + 11

    for c0, _ in bloques:
        letra = get_column_letter(c0 + 4)
        w.conditional_formatting.add(
            f"{letra}{T0}:{letra}{T1}",
            FormulaRule(formula=[f'AND({letra}{T0}<>"",{letra}{T0}>0)'], fill=verde))
        w.conditional_formatting.add(
            f"{letra}{T0}:{letra}{T1}",
            FormulaRule(formula=[f'AND({letra}{T0}<>"",{letra}{T0}<=0)'], fill=rojo))

    w.freeze_panes = "C6"
    w.cell(row=T1 + 2, column=2,
           value="«Mejora» ya tiene en cuenta el sentido de cada prueba: verde es progreso.").font = peque
    return w

hoja_test(
    "Test físico", "TESTS FÍSICOS",
    [("Salto vertical (cm)", True), ("Sprint 20 m (s)", False),
     ("Agilidad en T (s)", False), ("Resistencia (nivel)", True)],
    "Cuatro mediciones: primera semana de septiembre, 11 de enero, 22 de marzo y final de temporada."
)

hoja_test(
    "Test técnico", "TESTS TÉCNICOS",
    [("Tiros libres (/20)", True), ("Entradas 60\" (nº)", True),
     ("Circuito manejo (s)", False)],
    "Mismo calendario que los tests físicos. Comparte los resultados en privado con cada jugadora."
)

# =========================================================
# HOJA 6 — PANEL
# =========================================================
ws = wb.create_sheet("Panel")
ws.sheet_view.showGridLines = False
ws["A1"] = "PANEL DE TEMPORADA"
ws["A1"].font = titulo
ws["A2"] = "Todo se calcula solo. Revísalo al cerrar cada bloque: diciembre, marzo y abril."
ws["A2"].font = peque

ws["A4"] = "BALANCE"
ws["A4"].font = subtitulo

bal = [
    ("Partidos jugados", f"=COUNT(Partidos!E{R0}:E{R1})", "0"),
    ("Victorias", f"=COUNTIF(Partidos!H{R0}:H{R1},\"V\")", "0"),
    ("Derrotas", f"=COUNTIF(Partidos!H{R0}:H{R1},\"D\")", "0"),
    ("Puntos a favor (media)", f"=IFERROR(AVERAGE(Partidos!E{R0}:E{R1}),\"\")", "0.0"),
    ("Puntos en contra (media)", f"=IFERROR(AVERAGE(Partidos!F{R0}:F{R1}),\"\")", "0.0"),
    ("Diferencia media", f"=IFERROR(AVERAGE(Partidos!G{R0}:G{R1}),\"\")", "0.0"),
    ("Diferencia acumulada", f"=IFERROR(SUM(Partidos!G{R0}:G{R1}),\"\")", "0"),
]
r = 5
for et, f, nf in bal:
    a = ws.cell(row=r, column=1, value=et)
    a.font = normal
    a.border = borde
    b = ws.cell(row=r, column=2, value=f)
    b.font = bold
    b.fill = fill_calc
    b.alignment = ctr
    b.border = borde
    b.number_format = nf
    r += 1

r += 1
ws.cell(row=r, column=1, value="INDICADORES DEL MODELO DE JUEGO").font = subtitulo
r += 1
enc = ["Indicador", "Objetivo", "Media temporada", "Estado"]
for j, e in enumerate(enc):
    c = ws.cell(row=r, column=1 + j, value=e)
    c.font = cab
    c.fill = fill_cab
    c.alignment = ctr
    c.border = borde
r += 1
fila_ind = r

ind = [
    ("% rebote defensivo", 0.70, f"=IFERROR(AVERAGE(Partidos!K{R0}:K{R1}),\"\")", "0.0%", ">="),
    ("Puntos en transición", 20, f"=IFERROR(AVERAGE(Partidos!O{R0}:O{R1}),\"\")", "0.0", ">="),
    ("% rebote ofensivo", 0.30, f"=IFERROR(AVERAGE(Partidos!N{R0}:N{R1}),\"\")", "0.0%", ">="),
    ("Pérdidas", 18, f"=IFERROR(AVERAGE(Partidos!P{R0}:P{R1}),\"\")", "0.0", "<="),
]
for et, objv, f, nf, sentido in ind:
    ws.cell(row=r, column=1, value=et).font = normal
    o = ws.cell(row=r, column=2, value=objv)
    o.number_format = nf
    o.alignment = ctr
    o.font = normal
    v = ws.cell(row=r, column=3, value=f)
    v.number_format = nf
    v.alignment = ctr
    v.font = bold
    v.fill = fill_calc
    if sentido == ">=":
        est = f"=IF(C{r}=\"\",\"\",IF(C{r}>=B{r},\"CUMPLE\",\"POR DEBAJO\"))"
    else:
        est = f"=IF(C{r}=\"\",\"\",IF(C{r}<=B{r},\"CUMPLE\",\"POR ENCIMA\"))"
    e = ws.cell(row=r, column=4, value=est)
    e.alignment = ctr
    e.font = bold
    for j in range(1, 5):
        ws.cell(row=r, column=j).border = borde
    r += 1
fila_ind_fin = r - 1

ws.conditional_formatting.add(f"D{fila_ind}:D{fila_ind_fin}",
                              FormulaRule(formula=[f'D{fila_ind}="CUMPLE"'], fill=verde))
ws.conditional_formatting.add(f"D{fila_ind}:D{fila_ind_fin}",
                              FormulaRule(formula=[f'AND(D{fila_ind}<>"",D{fila_ind}<>"CUMPLE")'], fill=rojo))

r += 1
ws.cell(row=r, column=1, value="REPARTO DE MINUTOS").font = subtitulo
r += 1
colT = get_column_letter(col_ult + 1)
colM = get_column_letter(col_ult + 2)
colB = get_column_letter(col_ult + 4)
rep = [
    ("Minutos de la jugadora que más juega", f"=IFERROR(MAX(Minutos!{colT}{P0}:{colT}{P1}),\"\")"),
    ("Minutos de la jugadora que menos juega", f"=IFERROR(MIN(Minutos!{colT}{P0}:{colT}{P1}),\"\")"),
    ("Media de minutos del equipo", f"=IFERROR(AVERAGE(Minutos!{colM}{P0}:{colM}{P1}),\"\")"),
    ("Casos de menos de 10 min (todo el equipo)", f"=SUM(Minutos!{colB}{P0}:{colB}{P1})"),
]
for et, f in rep:
    a = ws.cell(row=r, column=1, value=et)
    a.font = normal
    a.border = borde
    b = ws.cell(row=r, column=2, value=f)
    b.font = bold
    b.fill = fill_calc
    b.alignment = ctr
    b.border = borde
    b.number_format = "0.0"
    r += 1

r += 2
ws.cell(row=r, column=1,
        value="Recordatorio: los partidos contra los equipos de la zona baja del grupo arrastran resultado a "
              "los grupos de descenso. Márcalos en la columna «Notas» de la hoja Partidos.").font = peque

ws.column_dimensions["A"].width = 42
ws.column_dimensions["B"].width = 18
ws.column_dimensions["C"].width = 18
ws.column_dimensions["D"].width = 16

# =========================================================
# CALENDARIO DE SESIONES
# =========================================================
festivos = {dt.date(2026,12,8),    # Inmaculada, martes
            dt.date(2026,12,25),   # Navidad, viernes
            dt.date(2027,1,1),     # Año Nuevo, viernes
            dt.date(2027,3,26)}    # Viernes Santo
voluntarias = {dt.date(2026,12,29), dt.date(2027,1,5), dt.date(2027,1,8)}

def bloque(d):
    if d <= dt.date(2026,10,2):  return "M0"
    if d <= dt.date(2026,12,18): return "M1"
    if d <= dt.date(2027,1,8):   return "Trans A"
    if d <= dt.date(2027,3,12):  return "M2"
    if d <= dt.date(2027,3,23):  return "Trans B"
    if d <= dt.date(2027,4,23):  return "M3"
    return "M4"

# martes 75 min · viernes 105 min
DUR = {1: 75, 4: 105}

sesiones = []
d = dt.date(2026, 9, 1)
fin = dt.date(2027, 5, 28)
while d <= fin:
    if d.weekday() in (1, 4) and d not in festivos:
        # parón navideño: solo las voluntarias
        if dt.date(2026,12,20) <= d <= dt.date(2027,1,8) and d not in voluntarias:
            d += dt.timedelta(days=1); continue
        sesiones.append((d, bloque(d), DUR[d.weekday()]))
    d += dt.timedelta(days=1)

DIAS = ["Lun", "Mar", "Mié", "Jue", "Vie", "Sáb", "Dom"]

def semana_de(d):
    lunes = d - dt.timedelta(days=d.weekday())
    return lunes

semanas = sorted({semana_de(x[0]) for x in sesiones})
idx_sem = {s: i + 1 for i, s in enumerate(semanas)}

# =========================================================
# HOJA — RPE
# =========================================================
ws = wb.create_sheet("RPE")
ws.sheet_view.showGridLines = False
ws["A1"] = "CONTROL DE CARGA · RPE"
ws["A1"].font = titulo
ws["A2"] = ("Al terminar la sesión, cada jugadora dice un número del 0 al 10. Anota la media del equipo. "
            "Referencia: 1-2 sesión de tiro estático · 10 intensidad de partido.")
ws["A2"].font = peque
ws["A3"] = ("Carga = RPE × minutos. Los martes duran 75 min y los viernes 105, así que las sesiones "
            "NO son comparables sin multiplicar. La columna Min viene rellena; cámbiala si varía.")
ws["A3"].font = peque

HDR = 5
cabs = [("Sem", 6), ("Fecha", 11), ("Día", 6), ("Bloque", 9), ("Min", 6),
        ("RPE\nmedio", 8), ("Carga\n(UA)", 9), ("Notas", 34)]
for j, (t, w) in enumerate(cabs, start=1):
    c = ws.cell(row=HDR, column=j, value=t)
    c.font = cab; c.fill = fill_cab; c.alignment = ctr_wrap; c.border = borde
    ws.column_dimensions[get_column_letter(j)].width = w
ws.row_dimensions[HDR].height = 26

R0 = HDR + 1
for i, (fecha, blq, dur) in enumerate(sesiones):
    r = R0 + i
    ws.cell(row=r, column=1, value=idx_sem[semana_de(fecha)]).alignment = ctr
    fc = ws.cell(row=r, column=2, value=fecha)
    fc.number_format = "DD/MM/YYYY"; fc.alignment = ctr
    ws.cell(row=r, column=3, value=DIAS[fecha.weekday()]).alignment = ctr
    ws.cell(row=r, column=4, value=blq).alignment = ctr
    mc = ws.cell(row=r, column=5, value=dur)
    mc.alignment = ctr; mc.font = entrada; mc.fill = fill_in
    rc = ws.cell(row=r, column=6)
    rc.alignment = ctr; rc.font = entrada; rc.fill = fill_in
    cc = ws.cell(row=r, column=7, value=f'=IF(F{r}="","",E{r}*F{r})')
    cc.alignment = ctr; cc.font = calc; cc.fill = fill_calc; cc.number_format = "0"
    nc = ws.cell(row=r, column=8); nc.font = entrada; nc.fill = fill_in
    for j in range(1, 5):
        ws.cell(row=r, column=j).font = normal
    for j in range(1, 9):
        ws.cell(row=r, column=j).border = borde
R1 = R0 + len(sesiones) - 1
ws.freeze_panes = "A6"

# resumen semanal a la derecha
c0 = 10
res = [("Sem", 6), ("Lunes de", 11), ("Carga semanal", 13), ("Δ vs anterior", 13), ("Aviso", 22)]
for j, (t, w) in enumerate(res):
    c = ws.cell(row=HDR, column=c0 + j, value=t)
    c.font = cab; c.fill = fill_cab; c.alignment = ctr_wrap; c.border = borde
    ws.column_dimensions[get_column_letter(c0 + j)].width = w

for i, s in enumerate(semanas):
    r = R0 + i
    n = i + 1
    ws.cell(row=r, column=c0, value=n).alignment = ctr
    fc = ws.cell(row=r, column=c0 + 1, value=s)
    fc.number_format = "DD/MM/YYYY"; fc.alignment = ctr
    ws.cell(row=r, column=c0 + 2,
            value=f'=IFERROR(IF(SUMIF($A${R0}:$A${R1},{c0 and get_column_letter(c0)}{r},$G${R0}:$G${R1})=0,"",'
                  f'SUMIF($A${R0}:$A${R1},{get_column_letter(c0)}{r},$G${R0}:$G${R1})),"")')
    prev = get_column_letter(c0 + 2) + str(r - 1)
    cur = get_column_letter(c0 + 2) + str(r)
    if i == 0:
        ws.cell(row=r, column=c0 + 3, value="")
    else:
        ws.cell(row=r, column=c0 + 3,
                value=f'=IFERROR(IF(OR({cur}="",{prev}=""),"",{cur}/{prev}-1),"")')
    dl = get_column_letter(c0 + 3) + str(r)
    ws.cell(row=r, column=c0 + 4,
            value=f'=IF({dl}="","",IF({dl}>0.2,"SUBIDA >20%",IF({dl}<-0.2,"Descarga","")))')
    for j in range(5):
        cc = ws.cell(row=r, column=c0 + j)
        cc.border = borde; cc.alignment = ctr
        if j >= 2:
            cc.font = bold; cc.fill = fill_calc
        else:
            cc.font = normal
    ws.cell(row=r, column=c0 + 2).number_format = "0"
    ws.cell(row=r, column=c0 + 3).number_format = "0%"
Rs1 = R0 + len(semanas) - 1

letra = get_column_letter(c0 + 4)
ws.conditional_formatting.add(f"{letra}{R0}:{letra}{Rs1}",
    FormulaRule(formula=[f'{letra}{R0}="SUBIDA >20%"'], fill=rojo))
ws.conditional_formatting.add(f"{letra}{R0}:{letra}{Rs1}",
    FormulaRule(formula=[f'{letra}{R0}="Descarga"'], fill=verde))

# =========================================================
# HOJA — ASISTENCIA
# =========================================================
ws = wb.create_sheet("Asistencia")
ws.sheet_view.showGridLines = False
ws["A1"] = "ASISTENCIA"
ws["A1"].font = titulo
ws["A2"] = "Escribe 1 si asiste y 0 si falta. Deja vacío si la sesión aún no se ha hecho."
ws["A2"].font = peque

HDR = 4
ws.cell(row=HDR, column=1, value="Fecha").font = cab
ws.cell(row=HDR, column=2, value="Bloque").font = cab
for j in (1, 2):
    c = ws.cell(row=HDR, column=j)
    c.fill = fill_cab; c.alignment = ctr; c.border = borde
ws.column_dimensions["A"].width = 11
ws.column_dimensions["B"].width = 9
for i in range(12):
    col = 3 + i
    c = ws.cell(row=HDR, column=col, value=f"J{i+1}")
    c.font = cab; c.fill = fill_cab; c.alignment = ctr; c.border = borde
    ws.column_dimensions[get_column_letter(col)].width = 6
c = ws.cell(row=HDR, column=15, value="Total")
c.font = cab; c.fill = fill_cab; c.alignment = ctr; c.border = borde
ws.column_dimensions["O"].width = 8

A0 = HDR + 1
for i, (fecha, blq, dur) in enumerate(sesiones):
    r = A0 + i
    fc = ws.cell(row=r, column=1, value=fecha)
    fc.number_format = "DD/MM/YYYY"; fc.alignment = ctr; fc.font = normal
    ws.cell(row=r, column=2, value=blq).alignment = ctr
    ws.cell(row=r, column=2).font = normal
    for j in range(3, 15):
        cc = ws.cell(row=r, column=j)
        cc.font = entrada; cc.fill = fill_in; cc.alignment = ctr
    tc = ws.cell(row=r, column=15, value=f"=IF(COUNT(C{r}:N{r})=0,\"\",SUM(C{r}:N{r}))")
    tc.font = bold; tc.fill = fill_calc; tc.alignment = ctr
    for j in range(1, 16):
        ws.cell(row=r, column=j).border = borde
A1 = A0 + len(sesiones) - 1

rt = A1 + 2
ws.cell(row=rt, column=2, value="Sesiones").font = bold
ws.cell(row=rt + 1, column=2, value="% asistencia").font = bold
for j in range(3, 15):
    col = get_column_letter(j)
    ws.cell(row=rt, column=j, value=f"=SUM({col}{A0}:{col}{A1})")
    ws.cell(row=rt + 1, column=j,
            value=f'=IFERROR(SUM({col}{A0}:{col}{A1})/COUNT({col}{A0}:{col}{A1}),"")')
    for k in (0, 1):
        cc = ws.cell(row=rt + k, column=j)
        cc.font = bold; cc.fill = fill_sec; cc.alignment = ctr; cc.border = borde
    ws.cell(row=rt + 1, column=j).number_format = "0%"
for k in (0, 1):
    for j in (1, 2):
        ws.cell(row=rt + k, column=j).fill = fill_sec
        ws.cell(row=rt + k, column=j).border = borde

ws.conditional_formatting.add(f"C{rt+1}:N{rt+1}",
    FormulaRule(formula=[f'AND(C{rt+1}<>"",C{rt+1}<0.8)'], fill=rojo))
ws.conditional_formatting.add(f"C{rt+1}:N{rt+1}",
    FormulaRule(formula=[f'AND(C{rt+1}<>"",C{rt+1}>=0.9)'], fill=verde))

ws.freeze_panes = "C5"
ws.cell(row=rt + 3, column=2,
        value="Sustituye J1…J12 por los nombres reales. Rojo por debajo del 80%.").font = peque

wb.save("seguimiento_temporada.xlsx")
print("OK ·", len(sesiones), "sesiones ·", len(semanas), "semanas")
